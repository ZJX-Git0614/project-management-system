from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).with_name("更新手册.docx")
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(92, 101, 112)
INK = RGBColor(30, 36, 44)
LIGHT_BLUE = "E8EEF5"
LIGHT_RED = "FDECEC"
LIGHT_GRAY = "F2F4F7"
DOCUMENT_FONT = "Microsoft YaHei"


def set_run_font(run, size=11, bold=False, color=INK, east_asia=DOCUMENT_FONT):
    run.font.name = east_asia
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), east_asia)
    r_fonts.set(qn("w:hAnsi"), east_asia)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    language = r_pr.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        r_pr.append(language)
    language.set(qn("w:val"), "zh-CN")
    language.set(qn("w:eastAsia"), "zh-CN")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text_run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = "1"
    text_run.append(text_node)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph._p.extend([begin, instruction, separate, text_run, end])
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, color=MUTED)


def add_body(doc, text, bold=False, color=INK, after=6):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.25
    set_run_font(paragraph.add_run(text), bold=bold, color=color)
    return paragraph


def create_decimal_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract.append(multi_level)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_format = OxmlElement("w:numFmt")
    num_format.set(qn("w:val"), "decimal")
    level.append(num_format)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    level.append(level_text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    level.append(suffix)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "271")
    p_pr.append(indent)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_numbered(doc, text, num_id):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.375)
    paragraph.paragraph_format.first_line_indent = Inches(-0.188)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])
    p_pr.insert(0, num_pr)
    set_run_font(paragraph.add_run(text))
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.375)
    paragraph.paragraph_format.first_line_indent = Inches(-0.188)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    set_run_font(paragraph.add_run(text))
    return paragraph


def add_code(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.right_indent = Inches(0.22)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.05
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GRAY)
    p_pr.append(shd)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = INK
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), DOCUMENT_FONT)
    return paragraph


def add_callout(doc, title, text, fill=LIGHT_BLUE, color=DARK_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    set_run_font(paragraph.add_run(title), bold=True, color=color)
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.2
    set_run_font(paragraph.add_run(text), color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.right_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = DOCUMENT_FONT
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25
normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), DOCUMENT_FONT)

for name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, DARK_BLUE, 10, 5),
):
    style = styles[name]
    style.font.name = DOCUMENT_FONT
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = color
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), DOCUMENT_FONT)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_run_font(header.add_run("CEASTAR PMS  |  WINDOWS OFFLINE UPDATE"), size=9, bold=True, color=MUTED)
add_page_field(section.footer.paragraphs[0])

title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(60)
title.paragraph_format.space_after = Pt(8)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(title.add_run("Ceastar项目管理系统"), size=28, bold=True, color=DARK_BLUE)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(18)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(subtitle.add_run("Windows x86 离线更新手册"), size=18, bold=True, color=BLUE)

meta = doc.add_table(rows=3, cols=2)
set_table_geometry(meta, [2200, 7160])
meta_data = [
    ("更新版本", "2026.07.27"),
    ("适用环境", "Windows 10/11 x86-64，已完成 Ceastar PMS 首次安装"),
    ("更新方式", "离线 Docker 镜像增量更新，保留数据库与上传文档"),
]
for row, (label, value) in zip(meta.rows, meta_data):
    shade_cell(row.cells[0], LIGHT_BLUE)
    set_run_font(row.cells[0].paragraphs[0].add_run(label), bold=True, color=DARK_BLUE)
    set_run_font(row.cells[1].paragraphs[0].add_run(value))

doc.add_paragraph().paragraph_format.space_after = Pt(12)
add_callout(
    doc,
    "核心原则",
    "更新脚本先备份数据库和上传文档，再替换应用容器。不要重新运行 install.bat，也不要执行 docker compose down -v。",
    fill=LIGHT_RED,
    color=RGBColor(155, 28, 28),
)

doc.add_page_break()

doc.add_heading("1. 更新包内容", level=1)
package_table = doc.add_table(rows=1, cols=2)
package_table.style = "Table Grid"
set_table_geometry(package_table, [2900, 6460])
for cell, text in zip(package_table.rows[0].cells, ("文件或目录", "用途")):
    shade_cell(cell, LIGHT_BLUE)
    set_run_font(cell.paragraphs[0].add_run(text), bold=True, color=DARK_BLUE)
for name, purpose in (
    ("update.bat / update.ps1", "自动校验、备份、加载镜像、迁移数据库并重建应用容器"),
    ("rollback.bat / rollback.ps1", "出现异常时恢复更新前保留的应用镜像"),
    ("images", "新版 linux/amd64 Docker 离线镜像"),
    ("image-name.txt", "新版镜像名称"),
    ("image.sha256", "镜像 SHA-256 完整性校验值"),
    ("更新手册", "Word 与纯文本操作说明"),
):
    cells = package_table.add_row().cells
    set_run_font(cells[0].paragraphs[0].add_run(name), bold=True)
    set_run_font(cells[1].paragraphs[0].add_run(purpose))
set_table_geometry(package_table, [2900, 6460])

doc.add_heading("2. 更新前准备", level=1)
prepare_numbering = create_decimal_numbering(doc)
add_numbered(doc, "启动 Docker Desktop，并等待其显示运行正常。", prepare_numbering)
add_numbered(doc, "确认系统盘或 Docker 数据盘至少有 4 GB 可用空间。", prepare_numbering)
add_numbered(doc, "等待当前用户操作保存完成，暂时关闭正在使用系统的浏览器页面。", prepare_numbering)
add_numbered(doc, "把整个更新包文件夹复制到原部署目录中，更新包文件夹与 docker-compose.yml 直接相邻。", prepare_numbering)
add_code(doc, "D:\\PMS\\Ceastar-PMS-内网完整部署-20260723\\docker-compose.yml\nD:\\PMS\\Ceastar-PMS-内网完整部署-20260723\\Ceastar-PMS-更新包-20260727\\update.bat")

doc.add_heading("3. 执行更新", level=1)
update_numbering = create_decimal_numbering(doc)
add_numbered(doc, "双击更新包中的 update.bat。", update_numbering)
add_numbered(doc, "保持命令窗口开启。脚本会校验镜像、备份数据、保留回退镜像并更新应用。", update_numbering)
add_numbered(doc, "窗口显示 Ceastar PMS update completed successfully. 后，按任意键关闭。", update_numbering)
add_numbered(doc, "打开 http://localhost:3000，使用原账号登录。", update_numbering)
add_callout(
    doc,
    "更新脚本不会修改",
    "PostgreSQL 数据卷、上传文档卷、Windows 上的 Ollama 安装与模型目录，以及已存在的备份文件。",
)

doc.add_heading("4. 更新后检查", level=1)
for item in (
    "能够打开登录页并使用原账号登录。",
    "项目、任务、事项、风险、预算和文档数据仍然存在。",
    "项目进度管理可以导入 MPP 文件，不再出现 path 参数为数字的错误。",
    "系统设置 > 系统数据管理可以查看备份记录。",
):
    add_bullet(doc, item)
add_body(doc, "如需检查容器状态，在原部署目录打开 PowerShell：", bold=True)
add_code(doc, "docker compose ps\ndocker compose logs --tail 100 pms")

doc.add_heading("5. 回退应用版本", level=1)
add_body(doc, "如果更新后系统无法正常使用，双击更新包中的 rollback.bat。")
add_body(doc, "回退脚本会先再次备份数据库和上传文档，再恢复更新前保留的应用镜像并重建 pms 容器。回退不会删除更新后的数据库数据。")

doc.add_heading("6. 重要警告", level=1)
for warning in (
    "不要运行 install.bat；该脚本仅用于首次安装，强制运行可能覆盖现有数据库。",
    "不要执行 docker compose down -v。",
    "不要删除 ceastar-pms_pgdata 和 ceastar-pms_document_storage 数据卷。",
    "不要只复制镜像 tar 文件，必须完整复制整个更新包。",
    "更新成功后，建议把原部署目录中新生成的 backups 文件夹复制到其他磁盘。",
):
    add_bullet(doc, warning)

doc.add_page_break()
doc.add_heading("7. 常见问题", level=1)
faq = doc.add_table(rows=1, cols=2)
faq.style = "Table Grid"
set_table_geometry(faq, [3200, 6160])
for cell, text in zip(faq.rows[0].cells, ("提示或现象", "处理方式")):
    shade_cell(cell, LIGHT_BLUE)
    set_run_font(cell.paragraphs[0].add_run(text), bold=True, color=DARK_BLUE)
for issue, solution in (
    ("Docker Desktop is not running", "启动 Docker Desktop，等待运行正常后重新执行 update.bat。"),
    ("checksum does not match", "镜像复制不完整，删除本次复制的更新包后重新复制。"),
    ("deployment directory was not found", "把更新包文件夹直接放入 docker-compose.yml 所在目录。"),
    ("页面暂时无法访问", "等待 1 至 3 分钟；仍失败时查看 pms 容器日志。"),
    ("Ollama 模型是否受影响", "不受影响，更新脚本不会修改 Windows 的 Ollama 模型目录。"),
):
    cells = faq.add_row().cells
    set_run_font(cells[0].paragraphs[0].add_run(issue), bold=True)
    set_run_font(cells[1].paragraphs[0].add_run(solution))
set_table_geometry(faq, [3200, 6160])

doc.core_properties.title = "Ceastar项目管理系统 Windows x86 离线更新手册"
doc.core_properties.subject = "Ceastar PMS 2026.07.27 离线更新操作说明"
doc.core_properties.author = "Ceastar项目管理系统"
doc.save(OUTPUT)
print(OUTPUT)
